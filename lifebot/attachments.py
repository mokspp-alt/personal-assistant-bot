"""Фото и файлы из Telegram → блоки для Claude: картинки, PDF, Excel, текстовые файлы."""

from __future__ import annotations

import base64
import csv
import io
from pathlib import PurePath

MAX_SIDE = 2576  # больше модели всё равно не используют
MAX_IMAGE_BYTES = 3_500_000  # API принимает до 5 МБ в base64, это ≈3,7 МБ файла
MAX_TEXT_CHARS = 200_000
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp", ".tif", ".tiff"}
TEXT_EXTENSIONS = {".txt", ".csv", ".tsv", ".json", ".md", ".log"}
EXCEL_EXTENSIONS = {".xlsx", ".xlsm"}


class UnsupportedFile(ValueError):
    pass


FORMATS = {"JPEG": "image/jpeg", "PNG": "image/png", "GIF": "image/gif", "WEBP": "image/webp"}


def _fit_image(data: bytes) -> tuple[bytes, str]:
    """Картинка для API: тип — по содержимому файла, а не по имени (иначе API отклонит запрос);
    большие уменьшаются до MAX_SIDE по длинной стороне и пережимаются в JPEG."""
    from PIL import Image

    with Image.open(io.BytesIO(data)) as img:
        media_type = FORMATS.get(img.format or "")
        if media_type and max(img.size) <= MAX_SIDE and len(data) <= MAX_IMAGE_BYTES:
            return data, media_type
        img.thumbnail((MAX_SIDE, MAX_SIDE))
        out = io.BytesIO()
        img.convert("RGB").save(out, "JPEG", quality=85)
        return out.getvalue(), "image/jpeg"


def _excel_to_csv(data: bytes) -> str:
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for sheet in book.worksheets:
        buf = io.StringIO()
        writer = csv.writer(buf)
        for row in sheet.iter_rows(values_only=True):
            if any(cell is not None for cell in row):
                writer.writerow(["" if cell is None else cell for cell in row])
        parts.append(f"# Лист «{sheet.title}»\n{buf.getvalue()}")
    return "\n".join(parts)


def _text_document(text: str, name: str) -> dict:
    if len(text) > MAX_TEXT_CHARS:
        text = text[:MAX_TEXT_CHARS] + "\n…(файл обрезан: дальше слишком длинно)"
    return {"type": "document", "title": name,
            "source": {"type": "text", "media_type": "text/plain", "data": text}}


def to_block(name: str, mime: str | None, data: bytes) -> dict:
    """Блок содержимого для Messages API. UnsupportedFile — если такой файл не понимаем."""
    ext = PurePath(name).suffix.lower()
    mime = (mime or "").lower()
    if mime.startswith("image/") or ext in IMAGE_EXTENSIONS:
        try:
            data, mime = _fit_image(data)
        except Exception as e:
            raise UnsupportedFile("Не получилось открыть картинку.") from e
        return {"type": "image", "source": {"type": "base64", "media_type": mime,
                                            "data": base64.b64encode(data).decode()}}
    if mime == "application/pdf" or ext == ".pdf":
        return {"type": "document", "title": name, "source": {
            "type": "base64", "media_type": "application/pdf", "data": base64.b64encode(data).decode()}}
    if ext in EXCEL_EXTENSIONS:
        try:
            return _text_document(_excel_to_csv(data), name)
        except Exception as e:
            raise UnsupportedFile("Не получилось прочитать таблицу Excel.") from e
    if mime.startswith("text/") or ext in TEXT_EXTENSIONS:
        return _text_document(data.decode("utf-8", errors="replace"), name)
    raise UnsupportedFile("Пока понимаю фото и картинки, PDF, таблицы Excel (.xlsx) и текстовые "
                          "файлы (.txt, .csv). Этот файл открыть не могу.")
